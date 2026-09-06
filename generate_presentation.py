import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Colors
COLOR_TEAL = RGBColor(0, 75, 92)
COLOR_MUSTARD = RGBColor(212, 161, 54)
COLOR_NAVY = RGBColor(0, 26, 51)
COLOR_WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE
    # Top-left teal corner
    left = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(1.5), Inches(1.5))
    left.fill.solid()
    left.fill.fore_color.rgb = COLOR_TEAL
    left.line.fill.background()
    # Bottom-right mustard
    right = slide.shapes.add_shape(1, Inches(8.5), Inches(6), Inches(1.5), Inches(1.5))
    right.fill.solid()
    right.fill.fore_color.rgb = COLOR_MUSTARD
    right.line.fill.background()

def create_slide(title, points):
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL
    p.font.name = "Montserrat"
    
    content_box = slide.shapes.add_textbox(Inches(1), Inches(1.2), Inches(8), Inches(5.8))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, point in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(point, tuple):
            p.text = point[0]
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = COLOR_MUSTARD
            p.space_before = Pt(8)
        else:
            p.text = point
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_NAVY
            p.level = 1 if point.startswith("  ") else 0
    return slide

# Generating 20 slides
slides_data = [
    ("Title Slide", [("Multilingual Historical Entity Linking", ""), ("Implementation, Enhancement & Comparative Analysis", "")]),
    ("Presentation Overview", ["1. PROJECT FOUNDATION", "2. IMPLEMENTATION STRATEGY", "3. ALTERNATIVE EXPLORATION", "4. COMPARATIVE ANALYSIS", "5. PROJECT EXECUTION", "6. LEARNING & FUTURE"]),
    ("Entity Linking Foundation", [("Concept", ""), "Match mentions to Wikipedia IDs", ("Mid-Sem Status", ""), "Baselines working", ("End-Sem Target", ""), "Evaluations complete"]),
    ("Project Motivation & Overview", [("WHY ENTITY LINKING?", ""), "Named entity disambiguation is critical in NLP", "Historical texts present unique challenges"]),
    ("Paper 1: MHEL-LLaMo - Overview & Selection", [("KEY CONTRIBUTIONS", ""), "Confidence-based routing approach", "Multilingual support (9 languages tested)"]),
    ("Why MHELL Was Our Successful Path", [("CORE INSIGHT", ""), "Confidence-Based Routing", ("SUCCESS FACTORS", ""), "Modular design", "Resource efficient"]),
    ("MHELL Implementation Timeline", [("WEEK 1-2", ""), "Foundation", ("WEEK 5-6", ""), "Optimization", ("WEEK 9-10", ""), "Evaluation"]),
    ("MHELL Implementation Details", [("MODULE 1: CANDIDATE GENERATION (BELA)", ""), ("MODULE 2: CONFIDENCE ROUTER (XGBoost)", "")]),
    ("Challenges Overcome in MHELL", [("CHALLENGE 1: API Compatibility", ""), ("CHALLENGE 3: GPU Memory (32GB)", "")]),
    ("Our Enhancements Beyond Base Paper", [("1. XGBoost Confidence Router", ""), "82% vs 70%", ("2. Error Analysis", "")]),
    ("Paper 2: mReFinED - Overview & Exploration", [("KEY CONTRIBUTIONS", ""), "Cross-encoder based ranking", ("WHY WE EXPLORED", "")]),
    ("mReFinED Implementation Attempt", [("APPROACH", ""), "Alternative pipeline: BELA → Cross-Encoder → Threshold"]),
    ("mReFinED: Unresolved Issues", [("CHALLENGE 1", ""), "Computational Scalability", ("CHALLENGE 2", ""), "Non-Latin Script Performance"]),
    ("Comparative Analysis: MHEL vs mReFinED", [("WINNER: MHELL in Most Categories", ""), "Performance: 50% vs 48% F1"]),
    ("10-Week Project Timeline & Milestones", [("WEEKS 1-2:", ""), "Foundation", ("WEEKS 9-10:", ""), "Finalization"]),
    ("Results & Observations", [("MHEL-LLaMo: COMPLETE", ""), "Key Metrics:", "XGBoost Routing: 82% accuracy (+12%)"]),
    ("Key Learning Outcomes", [("TECHNICAL INSIGHTS", ""), "Routing logic > LLM cost", ("PROJECT MANAGEMENT", "")]),
    ("Future Possibilities", [("IMMEDIATE (1-2 months)", ""), ("MID-TERM (2-4 months)", "")]),
    ("References 1", ["Primary References & Research Papers"]),
    ("References & Closing", ["Additional refs", ("PROJECT IMPACT", ""), "THANK YOU"])
]

for title, points in slides_data:
    create_slide(title, points)

img_dir = r"C:\Users\Dhruv\OneDrive\Desktop\Final Entity Linking\images"
img_mappings = [
    (7, "mhel_llamo_arch.png", Inches(4.5), Inches(2.0), Inches(4.5)),
    (15, "mewsli9_f1.png", Inches(4.5), Inches(2.0), Inches(4.5)),
    (2, "blink_benchmark_bar.png", Inches(5.0), Inches(2.0), Inches(4.0)),
    (3, "mvd_recall_curve.png", Inches(5.0), Inches(2.0), Inches(4.0))
]

for s_idx, fname, l, t, w in img_mappings:
    try:
        path = os.path.join(img_dir, fname)
        if os.path.exists(path):
            prs.slides[s_idx].shapes.add_picture(path, l, t, width=w)
    except:
        pass

out_path = r"C:\Users\Dhruv\OneDrive\Desktop\Final Entity Linking\Board_Review_Presentation.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")