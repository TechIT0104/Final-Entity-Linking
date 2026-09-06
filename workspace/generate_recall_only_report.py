from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import datetime


def shade_cell(cell, color):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    cell._element.get_or_add_tcPr().append(element)


def main() -> None:
    # Verified recall-only values
    paper_macro_recall = 58.8
    reproduced_macro_recall = 15.371511337909285 * 100
    reproduced_micro_recall = 24.99

    gap_points = paper_macro_recall - reproduced_macro_recall
    ratio = reproduced_macro_recall / paper_macro_recall if paper_macro_recall else 0.0

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("MEWSLI-9 Recall-Only Comparison", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.color.rgb = RGBColor(31, 78, 121)
    title_run.font.size = Pt(24)
    title_run.bold = True

    subtitle = doc.add_paragraph("Paper Metric Alignment Report (Recall Only)")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(13)

    date_line = doc.add_paragraph(f"Generated: {datetime.datetime.now().strftime('%B %d, %Y %H:%M')}")
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line.runs[0].font.size = Pt(10)
    date_line.runs[0].italic = True

    doc.add_paragraph()

    doc.add_heading("Metric Policy", level=1)
    doc.add_paragraph(
        "This report uses recall-only metrics for comparison with the paper. "
        "F1 is excluded from the main comparison table to keep the evaluation apples-to-apples."
    )

    doc.add_heading("Exact Recall Definitions Used", level=1)
    doc.add_paragraph("1. Paper Macro Recall: The paper-reported MEWSLI-9 macro-average recall (58.8%).")
    doc.add_paragraph(
        "2. Reproduced Macro Recall: Mean of language recalls from the MEWSLI-9 evaluator "
        "(Average recall line from server log)."
    )
    doc.add_paragraph(
        "3. Reproduced Micro Recall: tp/(tp+fn) as implemented in multilingual_e2e_evaluation_mewsli9.py "
        "and printed as 'Micro-avg'."
    )

    doc.add_heading("Recall-Only Comparison (Primary)", level=1)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Light Grid Accent 1"

    headers = ["Metric", "Value", "Source"]
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = value
        shade_cell(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    rows = [
        ("Paper Macro Recall", f"{paper_macro_recall:.2f}%", "Paper report"),
        (
            "Reproduced Macro Recall",
            f"{reproduced_macro_recall:.2f}%",
            "Server log: Average recall:0.15371511337909285",
        ),
        (
            "Reproduced Micro Recall",
            f"{reproduced_micro_recall:.2f}%",
            "Evaluator formula: tp/(tp+fn)",
        ),
    ]

    for row_idx, (metric, value, source) in enumerate(rows, start=1):
        row = table.rows[row_idx]
        row.cells[0].text = metric
        row.cells[1].text = value
        row.cells[2].text = source

    doc.add_paragraph()
    doc.add_heading("Gap Analysis", level=1)
    doc.add_paragraph(f"Absolute Gap (Paper - Reproduced Macro Recall): {gap_points:.2f} percentage points")
    doc.add_paragraph(f"Relative Match (Reproduced / Paper): {ratio:.3f}x")

    doc.add_heading("Evidence", level=1)
    doc.add_paragraph("MEWSLI-9 log path: /DATA/kmpooja/mrefined_option1/logs/mewsli9_final_20260416_232428.log")
    doc.add_paragraph("Recall lines extracted from log:")
    doc.add_paragraph("- Average recall:0.15371511337909285", style="List Bullet")
    doc.add_paragraph("- Micro-avg:24.99", style="List Bullet")

    doc.add_heading("Code-Level Confirmation", level=1)
    doc.add_paragraph(
        "The MEWSLI evaluator computes recall directly per language with metrics.get_recall() and then "
        "prints Average recall. It also prints Micro-avg using tp/(tp+fn), which is micro recall."
    )

    doc.add_heading("Current Overnight Execution", level=1)
    doc.add_paragraph("TR2016 run has been started in unattended background mode.")
    doc.add_paragraph("- Host: 172.20.70.80", style="List Bullet")
    doc.add_paragraph("- PID: 483192", style="List Bullet")
    doc.add_paragraph(
        "- Log: /DATA/kmpooja/mrefined_option1/logs/tr2016_full_20260419_232238.log",
        style="List Bullet",
    )

    output = r"C:\Users\Dhruv\OneDrive\Desktop\Entity Linking\MEWSLI9_Recall_Only_Comparison.docx"
    doc.save(output)
    print(f"Created: {output}")


if __name__ == "__main__":
    main()
